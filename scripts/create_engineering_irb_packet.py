"""Create the versioned ALGET engineering IRB revision packet as DOCX files."""
from __future__ import annotations

import json
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUTPUT = Path(r"C:\Users\jewoo\Desktop\_System\Work\25_Service\IRB\AlGET\2026_revision")
BLUE = "1F4E78"
LIGHT_BLUE = "DCE6F1"
AMBER = "FFF2CC"
GRAY = RGBColor(89, 89, 89)


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_width(cell, width_twips):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_twips))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_twips):
    """Freeze table/grid/cell widths so Word and LibreOffice agree."""
    table.autofit = False
    total_width = sum(widths_twips)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_width))
    tbl_w.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths_twips[index])


def mark_header_row(row):
    """Expose the first data-table row as a repeating accessibility header."""
    tr_pr = row._tr.get_or_add_trPr()
    header = tr_pr.find(qn("w:tblHeader"))
    if header is None:
        header = OxmlElement("w:tblHeader")
        tr_pr.append(header)
    header.set(qn("w:val"), "1")


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    p_pr.append(shading)


def configure(doc, short_title, status, body_size=10.25, body_space_after=5.5):
    section = doc.sections[0]
    # Keep the archival DOCX header/footer-free. LibreOffice maps DOCX
    # first/odd/even parts differently from Word and can push a header or page
    # field outside alternating pages. Document control is carried in the
    # front-matter table so removing those unstable parts loses no provenance.
    section.top_margin = Inches(.72)
    section.bottom_margin = Inches(.72)
    section.left_margin = Inches(.8)
    section.right_margin = Inches(.8)
    section.header_distance = Inches(.3)
    section.footer_distance = Inches(.3)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(body_size)
    normal.paragraph_format.space_after = Pt(body_space_after)
    normal.paragraph_format.line_spacing = 1.15
    for style_name, size in [("Title", 24), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(BLUE)
    styles["Heading 1"].paragraph_format.space_before = Pt(14)
    styles["Heading 1"].paragraph_format.space_after = Pt(6)

def clean_inline(text):
    text = re.sub(r"\[([^]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    return text.replace("**", "").replace("`", "")


def add_table(doc, rows, widths_twips=None):
    if not rows:
        return
    columns = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=columns)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    available = 9360
    widths = list(widths_twips) if widths_twips else [available // columns] * columns
    if len(widths) != columns or sum(widths) != available:
        raise ValueError(f"table widths must contain {columns} columns totaling {available} twips")
    set_table_geometry(table, widths)
    mark_header_row(table.rows[0])
    for r_index, row in enumerate(rows):
        for c_index in range(columns):
            cell = table.cell(r_index, c_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            value = clean_inline(row[c_index].strip()) if c_index < len(row) else ""
            cell.text = value
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for run in p.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8.5)
                    run.bold = r_index == 0
            if r_index == 0:
                shade(cell, LIGHT_BLUE)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_front_matter(doc, title, subtitle, status):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run("RESEARCH PROTOCOL REVISION")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(BLUE)
    p = doc.add_paragraph(style="Title")
    p.add_run(title)
    p.paragraph_format.space_after = Pt(5)
    p = doc.add_paragraph()
    run = p.add_run(subtitle)
    run.font.name = "Arial"
    run.font.size = Pt(12)
    run.font.color.rgb = GRAY
    p.paragraph_format.space_after = Pt(12)
    values = [
        ("Prepared", "August 25, 2026"),
        ("Primary course", "ALGET Bio-Inspired Design"),
        ("Study stage", status),
        ("Document control", "Version 2.3; supersedes the 2025 draft only after IRB amendment approval, approved-version reconciliation, and release sign-off"),
    ]
    for label, value in values:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(.08)
        p.paragraph_format.right_indent = Inches(.08)
        p.paragraph_format.space_after = Pt(1)
        shade_paragraph(p, LIGHT_BLUE)
        label_run = p.add_run(f"{label}: ")
        label_run.bold = True
        p.add_run(value)
    doc.add_paragraph()


def markdown_body(doc, markdown, skip_title=True):
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if not line:
            i += 1
            continue
        if skip_title and line.startswith("# "):
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|?\s*:?-+", lines[i + 1]):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                values = [v for v in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r"\s*:?-+:?\s*", v) for v in values):
                    rows.append(values)
                i += 1
            add_table(doc, rows)
            continue
        if line.startswith(">"):
            quoted_paragraphs = []
            current = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                quoted = lines[i].strip()[1:].strip()
                if quoted:
                    current.append(quoted)
                elif current:
                    quoted_paragraphs.append(" ".join(current))
                    current = []
                i += 1
            if current:
                quoted_paragraphs.append(" ".join(current))
            for quoted in quoted_paragraphs:
                p = doc.add_paragraph(clean_inline(quoted))
                p.paragraph_format.left_indent = Inches(.3)
                p.paragraph_format.right_indent = Inches(.3)
                p.paragraph_format.space_after = Pt(7)
                shade_paragraph(p, "F4F6F9")
            continue
        heading = re.match(r"^(#{2,4})\s+(.*)", line)
        if heading:
            level = min(len(heading.group(1)) - 1, 3)
            doc.add_heading(clean_inline(heading.group(2)), level=level)
            i += 1
            continue
        if line.startswith("- "):
            # Use an explicit hanging-indent bullet for the same reason as the
            # numbered paragraphs below: LibreOffice may merge adjacent Word
            # list-style paragraphs into a single visual line.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.42)
            p.paragraph_format.first_line_indent = Inches(-.30)
            p.paragraph_format.space_after = Pt(2)
            p.add_run("•   " + clean_inline(line[2:]))
            i += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)", line)
        if numbered:
            # Keep the source's explicit number. LibreOffice occasionally
            # merges adjacent custom-numbering paragraphs on odd pages, while
            # an explicit hanging-indent paragraph renders consistently in
            # Word, LibreOffice, and the archival PDF.
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(.45)
            p.paragraph_format.first_line_indent = Inches(-.30)
            p.paragraph_format.space_after = Pt(2)
            p.add_run(f"{numbered.group(1)}.   {clean_inline(numbered.group(2))}")
            i += 1
            continue
        paragraph = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(r"^(#{1,4})\s|^-\s|^\d+\.\s|^\|", lines[i]):
            paragraph.append(lines[i].strip())
            i += 1
        doc.add_paragraph(clean_inline(" ".join(paragraph)))


def save_markdown_doc(
    source,
    output_name,
    title,
    subtitle,
    short_title,
    status,
    body_size=10.25,
    body_space_after=5.5,
):
    doc = Document()
    configure(doc, short_title, status, body_size, body_space_after)
    add_front_matter(doc, title, subtitle, status)
    markdown_body(doc, source.read_text(encoding="utf-8"))
    doc.core_properties.title = title
    doc.core_properties.subject = "ALGET engineering intervention study IRB revision"
    doc.core_properties.author = "ALGET Research Team"
    doc.save(OUTPUT / output_name)


def save_assessment_doc():
    spec = json.loads((ROOT / "research/measures/bio_inspired_assessment_blueprint.json").read_text(encoding="utf-8"))
    doc = Document()
    configure(doc, "Assessment Blueprint", "Development only")
    add_front_matter(
        doc,
        "Bio-Inspired Design Assessment Blueprint",
        "Parallel pretest, posttest, and retention forms with validation gates",
        "Development only; do not administer as a confirmatory outcome",
    )
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(.08)
    p.paragraph_format.right_indent = Inches(.08)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    shade_paragraph(p, AMBER)
    run = p.add_run("CONTROLLED RETIREMENT NOTICE: Do not use the prior Assessment Items.docx. It includes a non-active course and multiple incorrect numerical answer keys. Retain it only as an archived draft.")
    run.bold = True
    doc.add_heading("Form architecture", level=1)
    add_table(doc, [
        ["Property", "Specification"],
        ["Forms", "A (pretest), B (posttest), C (retention)"],
        ["Items per form", str(spec["items_per_form"])],
        ["Selected response", str(spec["scoring"]["selected_response_items"])],
        ["Constructed transfer", str(spec["scoring"]["constructed_transfer_items"])],
        ["Target time", f"{spec['target_minutes_per_form']} minutes"],
        ["Confidence", "0 to 100 in 10-point increments"],
    ], widths_twips=[3200, 6160])
    doc.add_heading("Content blueprint", level=1)
    rows = [["Domain", "Modules", "Selected/form", "Constructed/form"]]
    for domain in spec["content_domains"]:
        readable_domain = domain["domain"].replace("_and_", ", ").replace("_", " ").capitalize()
        rows.append([readable_domain, ", ".join(domain["modules"]), str(domain["selected_items_per_form"]), str(domain["constructed_items_per_form"])])
    add_table(doc, rows, widths_twips=[3000, 3000, 1680, 1680])
    doc.add_heading("Parallel-form construction rules", level=1)
    for rule in spec["parallel_form_rules"]:
        doc.add_paragraph(" " + rule, style="List Bullet")
    doc.add_heading("Required validation", level=1)
    readable_validation = {
        "subject_matter_experts": "6 to 8 experts",
        "measurement_experts": "3 to 5 experts",
        "cognitive_interviews": "8 to 12 students",
        "measurement_pilot": "30 to 50 students",
        "confirmatory_release": "All accuracy flags resolved; form statistics reviewed",
    }
    add_table(
        doc,
        [["Gate", "Minimum evidence"]]
        + [[key.replace("_", " ").title(), readable_validation[key]] for key in spec["validation_gates"]],
        widths_twips=[3600, 5760],
    )
    doc.add_heading("Constructed-response rubric", level=1)
    doc.add_paragraph("Score each dimension from 0 to 3. Write behavioral anchors and exemplars during expert review, then freeze them before the pilot.")
    add_table(doc, [
        ["Dimension", "0", "1", "2", "3"],
        ["Mechanism accuracy", "Absent/incorrect", "Partly correct", "Correct with minor gap", "Accurate and causally explicit"],
        ["Evidence alignment", "No evidence", "Evidence named", "Relevant evidence linked", "Multiple relevant cues integrated"],
        ["Constraint reasoning", "No constraint", "Constraint listed", "Trade-off explained", "Trade-off evaluated with decision"],
        ["Transfer justification", "Surface analogy", "Feature match", "Mechanism mapped", "Mechanism and boundary conditions mapped"],
    ])
    doc.add_heading("Item-quality control", level=1)
    doc.add_paragraph("Every numerical item requires an independent derivation sheet, unit check, distractor rationale, and second-person key verification. Every item must map to the active content version and a misconception sidecar. Record expert CVI ratings, cognitive-interview findings, pilot difficulty/discrimination, response time, missingness, and revision history.")
    doc.core_properties.title = "ALGET Bio-Inspired Design Assessment Blueprint"
    doc.core_properties.author = "ALGET Research Team"
    doc.save(OUTPUT / "03_BioInspired_Assessment_Blueprint_v2.docx")


def main():
    OUTPUT.mkdir(parents=True, exist_ok=True)
    save_markdown_doc(
        ROOT / "docs/ENGINEERING_INTERVENTION_STUDY_PROTOCOL.md",
        "01_Engineering_Intervention_Protocol_v2.docx",
        "ALGET Engineering Intervention Study",
        "Protocol, analysis plan, sample-size justification, and go/no-go gates",
        "Engineering Intervention Protocol",
        "Proposed expanded-study implementation packet; IRB amendment or authoritative expanded-scope approval required before recruitment",
        body_size=10,
        body_space_after=5,
    )
    save_markdown_doc(
        ROOT / "research/measures/engineering_study_instruments.md",
        "02_Engineering_Study_Instruments_v2.docx",
        "ALGET Engineering Study Instruments",
        "Administration schedule, scoring codebook, and interview guide",
        "Engineering Study Instruments",
        "Proposed expanded-study review attachment; not authorized for enrolled-student administration",
        body_size=10,
        body_space_after=5,
    )
    save_assessment_doc()
    save_markdown_doc(
        ROOT / "docs/ENGINEERING_RECRUITMENT_AND_CONSENT_DRAFT.md",
        "04_Engineering_Recruitment_Consent_Draft_v2.docx",
        "ALGET Engineering Study Participant Consent",
        "Recruitment, electronic consent, privacy, compensation, and participant-rights language",
        "Engineering Participant Consent",
        "Proposed expanded-study working copy; obtain and match the stamped amended consent before recruitment",
        body_size=10,
        body_space_after=5,
    )
    print(f"created 4 DOCX files in {OUTPUT}")


if __name__ == "__main__":
    main()
