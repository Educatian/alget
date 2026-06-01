from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt


OUT_DIR = Path(__file__).resolve().parent
PLATFORM_URL = "https://significant-upgrade.alget.pages.dev"


COURSES = {
    "CAT100": {
        "file": "CAT100_ALGET_Announcement_Draft.docx",
        "title": "CAT 100 Announcement Draft: ALGET Supplemental Learning Guide",
        "course": "CAT 100: Computer Applications / Digital Skills",
        "audience": "CAT 100 Summer I and CAT 100 Summer II students",
        "guide_url": "https://significant-upgrade.alget.pages.dev/guides/cat100/index.html",
        "subject": "Use ALGET as a supplemental practice guide for CAT 100",
        "opening": (
            "Dear CAT 100 students, this course includes access to ALGET, a "
            "supplemental learning space designed to help you practice, review, "
            "and connect the digital skills we work on in class."
        ),
        "purpose": [
            (
                "ALGET is a guided practice and review space. It is meant to help "
                "you revisit course ideas, try short practice tasks, and build "
                "confidence with tools such as productivity software, data work, "
                "digital portfolios, AI-supported teacher communication, and web "
                "publishing."
            ),
            (
                "ALGET does not replace the official course site, assignment "
                "instructions, due dates, gradebook, instructor feedback, or class "
                "announcements. Always follow the instructions posted in the LMS "
                "and shared by your instructor."
            ),
        ],
        "entry_steps": [
            f"Open ALGET: {PLATFORM_URL}",
            'Click "Sign in to start."',
            'Under "Current students," type your real first and last name.',
            'Choose the correct cohort: "CAT 100 Summer I" or "CAT 100 Summer II."',
            'Click "Enter my course."',
        ],
        "weekly_use": [
            (
                "Use the CAT 100 guide to see the recommended pathway for the "
                "course: digital citizenship, resume and professional branding, "
                "Excel and data stories, AI-supported teacher communication, "
                "presentations, GitHub Pages, productivity routines, and final "
                "transfer work."
            ),
            (
                "After each class topic, use the matching ALGET section for extra "
                "practice. If a diagnostic or short check appears, complete it so "
                "the system can recommend useful review."
            ),
            (
                "Use highlights, notes, and help prompts when something is unclear. "
                "These tools are for learning support, not for submitting private "
                "or sensitive information."
            ),
        ],
        "privacy": (
            "Please enter your real name so your learning activity can be connected "
            "to the correct student. ALGET may record learning interactions such as "
            "page progress, practice attempts, highlights, notes, and help requests "
            "to support learning, troubleshooting, and course improvement. Do not "
            "enter private or sensitive personal information in notes or chat."
        ),
        "help": (
            "If you cannot access your course, selected the wrong cohort, or do not "
            "see the expected CAT 100 materials, contact your instructor with your "
            "full name, cohort, and a screenshot of what you see."
        ),
        "short": (
            "CAT 100 students: ALGET is available as a supplemental guided practice "
            f"space for this course. Open {PLATFORM_URL}, click \"Sign in to start,\" "
            "enter your real name under Current students, and choose the correct "
            "CAT 100 Summer I or Summer II cohort. Use the CAT 100 guide here: "
            "https://significant-upgrade.alget.pages.dev/guides/cat100/index.html. "
            "ALGET supports review and practice, but the LMS remains the official "
            "source for assignments, due dates, grades, and instructor feedback."
        ),
    },
    "CAT531": {
        "file": "CAT531_ALGET_Announcement_Draft.docx",
        "title": "CAT 531 Announcement Draft: ALGET Supplemental Learning Guide",
        "course": "CAT 531: Educational Technology / AI-Supported Teaching Practice",
        "audience": "CAT 531 students",
        "guide_url": "https://significant-upgrade.alget.pages.dev/guides/cat531/index.html",
        "subject": "Use ALGET as a supplemental teaching-technology practice guide for CAT 531",
        "opening": (
            "Dear CAT 531 students, this course includes access to ALGET, a "
            "supplemental learning space designed to help you connect educational "
            "technology concepts with practical teaching decisions."
        ),
        "purpose": [
            (
                "ALGET is a guided practice and reflection space. It is meant to "
                "help you think through instructional design choices, AI-supported "
                "teaching tools, digital teaching scenarios, policy questions, "
                "technology evaluation, portfolio evidence, and your professional "
                "vision as an educator."
            ),
            (
                "ALGET is not a replacement for the official LMS, readings, "
                "assignment instructions, due dates, gradebook, instructor feedback, "
                "or class announcements. Treat it as a supplement that helps you "
                "prepare, review, and deepen your thinking."
            ),
        ],
        "entry_steps": [
            f"Open ALGET: {PLATFORM_URL}",
            'Click "Sign in to start."',
            'Under "Current students," type your real first and last name.',
            'Choose the "CAT 531" cohort.',
            'Click "Enter my course."',
        ],
        "weekly_use": [
            (
                "Use the CAT 531 guide to follow the recommended course pathway: "
                "digital teaching scenarios, TeachGen@i, Ethobot, AI policy, "
                "edtech evaluation, final professional vision, and portfolio "
                "evidence."
            ),
            (
                "Before or after each course topic, use the matching ALGET section "
                "for guided review, short practice, and reflection. The strongest "
                "use of the system is not simply finding an answer; it is explaining "
                "why a tool, policy, or design decision would make sense in a "
                "teaching context."
            ),
            (
                "Use highlights, notes, and help prompts to capture questions, "
                "teaching connections, or places where you want more clarification."
            ),
        ],
        "privacy": (
            "Please enter your real name so your learning activity can be connected "
            "to the correct student. ALGET may record learning interactions such as "
            "page progress, practice attempts, highlights, notes, and help requests "
            "to support learning, troubleshooting, and course improvement. Do not "
            "enter private student information, school records, or sensitive personal "
            "information in notes or chat."
        ),
        "help": (
            "If you cannot access CAT 531 materials or something does not look "
            "right, contact your instructor with your full name and a screenshot of "
            "the issue."
        ),
        "short": (
            "CAT 531 students: ALGET is available as a supplemental guided practice "
            f"and reflection space for this course. Open {PLATFORM_URL}, click "
            "\"Sign in to start,\" enter your real name under Current students, and "
            "choose CAT 531. Use the CAT 531 guide here: "
            "https://significant-upgrade.alget.pages.dev/guides/cat531/index.html. "
            "ALGET supports review, reflection, and teaching-technology practice, "
            "but the LMS remains the official source for assignments, due dates, "
            "grades, and instructor feedback."
        ),
    },
}


def set_base_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    for style_name, size in [
        ("Title", 24),
        ("Heading 1", 18),
        ("Heading 2", 14),
        ("Heading 3", 12),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)


def add_meta(document: Document, label: str, value: str) -> None:
    paragraph = document.add_paragraph()
    run = paragraph.add_run(f"{label}: ")
    run.bold = True
    paragraph.add_run(value)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.add_run(item)


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        paragraph = document.add_paragraph(style="List Number")
        paragraph.add_run(item)


def build_doc(course_id: str, data: dict[str, object]) -> Path:
    document = Document()
    set_base_styles(document)

    section = document.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run(data["title"])

    note = document.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.add_run("Draft for instructor editing and LMS posting").italic = True

    add_meta(document, "Course", data["course"])
    add_meta(document, "Audience", data["audience"])
    add_meta(document, "Suggested posting time", "Before students begin the first ALGET activity")
    add_meta(document, "Platform link", PLATFORM_URL)
    add_meta(document, "Student guide link", data["guide_url"])

    document.add_heading("Subject Line", level=1)
    document.add_paragraph(data["subject"])

    document.add_heading("Announcement Text", level=1)
    document.add_paragraph(data["opening"])

    document.add_heading("What ALGET Is", level=2)
    for paragraph in data["purpose"]:
        document.add_paragraph(paragraph)

    document.add_heading("How to Enter the Course", level=2)
    add_numbered(document, data["entry_steps"])

    document.add_heading("How to Use It During the Course", level=2)
    add_bullets(document, data["weekly_use"])

    document.add_heading("Important Boundaries", level=2)
    add_bullets(
        document,
        [
            "Use the LMS as the official source for assignments, due dates, grades, and course announcements.",
            "Use ALGET as supplemental practice, review, reflection, and study support.",
            "If ALGET and the LMS ever appear to disagree, follow the LMS and ask your instructor.",
        ],
    )

    document.add_heading("Learning Data and Privacy Note", level=2)
    document.add_paragraph(data["privacy"])

    document.add_heading("Where to Get Help", level=2)
    document.add_paragraph(data["help"])

    document.add_heading("Optional Short Version", level=1)
    document.add_paragraph(data["short"])

    output_path = OUT_DIR / data["file"]
    document.save(output_path)
    return output_path


def main() -> None:
    for course_id, data in COURSES.items():
        path = build_doc(course_id, data)
        print(path)


if __name__ == "__main__":
    main()
